"""
Parser for the DCS codes Excel file.

Adapted from DCS Tool/dcs_parser.py — ApplicationConfig dependency replaced
with a plain excel_path argument so it can be used standalone or from the
init script without any external config system.
"""
import logging
import re
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook
from pandas import DataFrame

logger = logging.getLogger(__name__)

# Matches 1-2 non-alphanumeric characters only (expansion indicator in column 2)
_EXPANSION_INDICATOR_RE = re.compile(r'^[^a-zA-Z0-9]{1,2}$')

# Matches a replacement row: 1-2 non-alphanumeric chars, then whitespace, then values
_REPLACEMENT_ROW_RE = re.compile(r'^([^a-zA-Z0-9]{1,2})\s+(.+)$')


def _pad_to_3(val) -> str:
    """Return val as exactly 3 characters, left-justified, space-padded. Never strips."""
    import pandas as pd
    if pd.isna(val):
        return "   "
    return str(val).ljust(3)[:3]


class DcsGsParser:
    """Parser for DCS Code data in an Excel file. Normalizes data and expands
    all expansion indicators (asterisks etc.) into full rows."""

    def __init__(self, excel_path: str):
        self.excel_path = excel_path
        self.initial_df = pd.read_excel(excel_path, sheet_name=0, header=0)
        self.result_df = None
        self._cleaned_df = self._clear_header_rows(self.initial_df)
        self._parse_it_all()

    def _clear_header_rows(self, df) -> DataFrame:
        mask = df.iloc[:, 0].astype(str).str.contains("DCS Codes", na=False)
        df.loc[mask] = pd.NA
        return df

    def _parse_it_all(self):
        narrow_df = self._wide_to_narrow(self._cleaned_df)
        result = self._reduce_multiple_empties_rows_to_one(narrow_df)
        blocks = self._parse_into_blocks(result)

        final_df = pd.DataFrame()
        for block in blocks:
            block_df = self._parse_block(block)
            final_df = pd.concat([final_df, block_df], ignore_index=True)

        final_df.columns = ["D", "C", "S", "Description", "Type", "DCS"]
        final_df = final_df[["DCS", "D", "C", "S", "Type", "Description"]]
        for col in ["D", "C", "S"]:
            final_df[col] = final_df[col].apply(_pad_to_3)
        final_df["DCS"] = final_df["DCS"].str.strip()
        final_df = final_df.sort_values(by="DCS")
        self.result_df = final_df

    def _wide_to_narrow(self, df) -> DataFrame:
        blocks = []
        total_cols = len(df.columns)
        for start_col in range(0, total_cols, 5):
            if start_col + 4 <= total_cols:
                block = df.iloc[:, start_col:start_col + 4].copy()
                block.columns = range(4)
                block = pd.concat(
                    [block, pd.DataFrame([[None] * 4], columns=range(4))],
                    ignore_index=True,
                )
                blocks.append(block)
        return pd.concat(blocks, axis=0, ignore_index=True) if blocks else df

    def _reduce_multiple_empties_rows_to_one(self, df) -> DataFrame:
        if df is None or df.empty:
            return df
        empty_rows_mask = df.isna().all(axis=1)
        rows_to_keep = []
        in_empty_sequence = False
        for idx in range(len(df)):
            is_empty = empty_rows_mask.iloc[idx]
            if is_empty:
                if not in_empty_sequence:
                    rows_to_keep.append(idx)
                    in_empty_sequence = True
            else:
                rows_to_keep.append(idx)
                in_empty_sequence = False
        return df.iloc[rows_to_keep].reset_index(drop=True)

    def _parse_into_blocks(self, df) -> list[DataFrame]:
        blocks = []
        current_block = []
        for idx, row in df.iterrows():
            if row.isna().all():
                if current_block:
                    blocks.append(pd.DataFrame(current_block))
                    current_block = []
            else:
                current_block.append(row)
        return blocks

    def _parse_block(self, block: DataFrame) -> DataFrame:
        if block is None or block.empty:
            return block

        df = block.copy()
        df = df.dropna(how='all').reset_index(drop=True)
        if df.empty:
            return df

        # Extract title from first row
        title = str(df.iloc[0, 0]) if pd.notna(df.iloc[0, 0]) else ""
        df = df.iloc[1:].reset_index(drop=True)
        if df.empty:
            return df

        # Collect replacement rows from the bottom
        replacements: dict[str, list[str]] = {}
        rows_to_drop = []
        for i in range(len(df) - 1, -1, -1):
            first_col = str(df.iloc[i, 0]) if pd.notna(df.iloc[i, 0]) else ""
            first_col = first_col.strip()
            m = _REPLACEMENT_ROW_RE.match(first_col)
            if m:
                symbol = m.group(1)
                values = [v.strip() for v in m.group(2).split(',') if v.strip()]
                replacements[symbol] = values
                rows_to_drop.append(i)
            else:
                break

        if rows_to_drop:
            df = df.drop(df.index[rows_to_drop]).reset_index(drop=True)
        if df.empty:
            return df

        # Forward-fill columns 0, 1, 3
        for col_idx in [0, 1, 3]:
            if col_idx < len(df.columns):
                df.iloc[:, col_idx] = df.iloc[:, col_idx].replace('', pd.NA).ffill()

        # Expand expansion indicators in column 2
        expanded_rows = []
        for idx, row in df.iterrows():
            col2_value = str(row.iloc[2]) if pd.notna(row.iloc[2]) else ""
            col2_value = col2_value.strip()
            if _EXPANSION_INDICATOR_RE.match(col2_value):
                if col2_value not in replacements:
                    logger.warning(
                        f"Expansion indicator '{col2_value}' at row {idx} has no replacement list — keeping as-is."
                    )
                    expanded_rows.append(row)
                else:
                    for replacement in replacements[col2_value]:
                        new_row = row.copy()
                        new_row.iloc[2] = replacement
                        expanded_rows.append(new_row)
            else:
                expanded_rows.append(row)

        if expanded_rows:
            df = pd.DataFrame(expanded_rows).reset_index(drop=True)

        df['Type'] = title

        df = df.fillna('')
        df['DCS'] = (
            df.iloc[:, 0].apply(_pad_to_3)
            + df.iloc[:, 1].apply(_pad_to_3)
            + df.iloc[:, 2].apply(_pad_to_3)
        )
        return df

    def write_normalized_sheet(self):
        """Write result_df to a 'Normalized' worksheet in the source Excel file."""
        if self.result_df is None or self.result_df.empty:
            logger.warning("No data to write — result_df is empty.")
            return
        wb = load_workbook(self.excel_path)
        if 'Normalized' in wb.sheetnames:
            del wb['Normalized']
        ws = wb.create_sheet('Normalized')
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws['A1'] = f"Generated: {timestamp}"
        for col, header in enumerate(self.result_df.columns, 1):
            ws.cell(row=2, column=col, value=header)
        for row_idx, row in enumerate(self.result_df.fillna("").astype(str).values.tolist(), 3):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        wb.save(self.excel_path)
        logger.info(f"Wrote {len(self.result_df)} rows to 'Normalized' sheet in {self.excel_path}.")
